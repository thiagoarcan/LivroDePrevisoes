from pathlib import Path
from docx import Document
import fitz
import hashlib
import json

root = Path(__file__).resolve().parents[2]
out = root / 'auditoria_interface'
doc = Document(root / 'Adaptação da aplicação para tablets e celulares.docx')
lines = []
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        lines.append(f'P{i:04d} [{p.style.name}] {p.text}')
for ti, t in enumerate(doc.tables):
    lines.append(f'TABELA {ti}')
    for ri, r in enumerate(t.rows):
        lines.append(f'T{ti}R{ri}: ' + ' | '.join(c.text.replace('\n', ' / ') for c in r.cells))
(out / 'dados/requisitos_extraidos.txt').write_text('\n'.join(lines), encoding='utf-8')
meta = []
for file in root.glob('Designer*.pdf'):
    d = fitz.open(file)
    info = {'file': file.name, 'sha256': hashlib.sha256(file.read_bytes()).hexdigest(),
            'pages': len(d), 'page_sizes': []}
    texts = []
    for n, page in enumerate(d):
        info['page_sizes'].append([page.rect.width, page.rect.height])
        texts.append(page.get_text())
        scale = min(1.5, 1600 / page.rect.width)
        page.get_pixmap(matrix=fitz.Matrix(scale, scale)).save(
            out / 'referencias' / f'{file.stem}_pagina_{n+1:02d}.png')
    (out / 'dados' / f'{file.stem}_texto.txt').write_text('\n'.join(texts), encoding='utf-8')
    meta.append(info)
(out / 'dados/referencias_metadados.json').write_text(
    json.dumps(meta, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(meta, ensure_ascii=False, indent=2))
print('DOCX_PARAGRAPHS', len(doc.paragraphs), 'TABLES', len(doc.tables), 'EXTRACTED_LINES', len(lines))
