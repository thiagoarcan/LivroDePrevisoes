"""Consolida evidências em relatório HTML, Markdown, Word e matrizes CSV."""
from pathlib import Path
from collections import Counter
from datetime import datetime
import csv, hashlib, html, json, re
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'auditoria_interface'
DATA=OUT/'dados'
def read(name):return json.loads((DATA/name).read_text(encoding='utf-8'))
def write(name,obj):(DATA/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')

F=[]
def finding(id,priority,title,expected,observed,cause,evidence,acceptance,kind='Confirmado em execução'):
    F.append(dict(id=id,prioridade=priority,titulo=title,esperado=expected,observado=observed,
                  causa=cause,evidencia=evidence,aceite=acceptance,tipo=kind))

finding('GUI-01','P1','Modo automático não aplica a composição de celular e tablet',
 'DOCX §§1.1–1.4: reorganizar automaticamente por espaço disponível, inclusive navegação inferior e lista/detalhe.',
 'Em 390px o tier é mobile, porém data-layout-mode está ausente; app-nav continua flex e bottom nav none. Em 820px ocorre o equivalente para tablet.',
 'applyResponsiveLayoutMode e handleViewportResize atribuem o tier automático somente para widescreen e ultrawide; CSS mobile/tablet exige o atributo.',
 'parts_js_ui.py:95,108,168; parts_css.py:1665,1823; metricas_navegador.json; celular_auto_viewDashboard.png; tablet_auto_viewDashboard.png',
 'Auto deve ativar a composição adequada em 320–1199px, inclusive após redimensionamento, sem depender de seleção manual.')
finding('GUI-02','P1','Cabeçalho e controles extrapolam a largura disponível',
 'DOCX P0012, §§2.2 e20: nenhuma rolagem horizontal da página inteira; cabeçalho compacto e console acessível.',
 '390px Auto: scrollWidth1077px;390px Celular:459–472px;820px Auto:1077px;1024px Tablet:1130px. O cabeçalho móvel manual mede cerca de227px de altura e mantém ações fora da tela.',
 'Grupos flex e seletor de visualização preservam larguras/conteúdo de desktop; a composição não comprime ou redistribui todos os grupos.',
 'parts_css.py:349,1277,1681; metricas_navegador.json; celular_manual_viewDashboard.png',
 'documentElement.scrollWidth deve ser igual à largura útil em todos os modos viáveis; console, sino, menu e nova previsão acessíveis sem deslocar a página inteira.')
finding('GUI-03','P1','Barra inferior móvel usa classes incompatíveis com o CSS',
 'DOCX §2.3 e Designer.pdf: quatro destinos com ícone/rótulo e apenas um comando Novo em posição previsível.',
 'No manual Celular a barra aparece, mas seus botões têm aparência nativa branca e o Novo é um quadrado pequeno; no Auto nem aparece.',
 'HTML usa bottom-nav-item, bottom-nav-fab e bottom-nav-badge; CSS estiliza bottom-nav-btn, mobile-fab e bottom-badge.',
 'parts_html.py:1628; parts_css.py:1574,1608; celular_manual_viewDashboard.png; celular_manual_viewNewEvent.png',
 'Unificar seletores; quatro destinos com estados ativo/foco, alvos de toque adequados, safe-area e sem botão concorrente.')
finding('GUI-04','P1','Faixa persistente de alarmes fica atrás do cabeçalho',
 'DOCX §3.2: faixa sempre visível logo abaixo do cabeçalho enquanto houver triggered.',
 'Com dois alarmes e rolagem, cabeçalho ocupa y0–227px; faixa fica em y52–152px, dentro da área ocupada pelo cabeçalho. Há sobreposição objetiva.',
 'alarm-banner usa top:52px fixo, apesar de cabeçalho móvel variável e camada superior.',
 'parts_css.py:1257; testes_estados.json:banner_scrolled; estado_banner_rolagem.png',
 'Deslocamento da faixa deve acompanhar a altura real do cabeçalho; reconhecimento e quantidade devem permanecer visíveis em retrato e paisagem.')
finding('GUI-05','P1','Filtros móveis não abrem como painel inferior',
 'DOCX §4.1 e Designer.pdf: painel inferior com rolagem própria, Limpar/Aplicar e contagem de filtros.',
 'Execução: class toolbar-filters,position static,height aproximadamente665px; o painel permanece no fluxo e pode ultrapassar a altura útil.',
 'CSS depende de #filterDrawer.active-sheet; JS altera somente style.display e não adiciona active-sheet.',
 'parts_js_ui.py:1530; parts_css.py:1727; testes_fluxos.json:filter_sheet; celular_manual_modal_filtros.png',
 'Abrir sheet limitado ao viewport, manter rodapé acessível, foco contido e fechamento explícito; considerar Apenas Console Ativo na contagem.')
finding('GUI-06','P1','Spotlight mistura descrição de um evento com tempo de outro',
 'DOCX §§1.3,3.3,14: descrição, hierarquia, previsto e contador devem corresponder ao mesmo evento do console ativo.',
 'Após pesquisar PIG em Pendentes e voltar ao Painel, a descrição continua na bomba atrasada12min, mas updateLiveTimers troca o contador para Faltam2h29min, referente ao PIG.',
 'renderDashboard seleciona spotlightTarget; updateLiveTimers usa activePendingEvents[0], uma lista filtrada e independente.',
 'parts_js_ui.py:1738,2246; testes_fluxos.json:spotlight_before/after; dashboard_contador_evento_incorreto.png',
 'Persistir ID/timestamp do destaque e usar a mesma origem em renderização e atualização; filtros da lista não devem alterar seu contador.')
finding('GUI-07','P1','Contador do modal de alarme fica congelado',
 'DOCX §7 e Designer.pdf: alarme exibe contador temporal destacado e atual.',
 'Avanço simulado de90s: texto do modal permanece igual, enquanto formatTimeRemaining já retorna outro atraso.',
 'alarmModalCountdown é preenchido apenas em showAlarmModal; updateLiveTimers não o inclui.',
 'parts_js_ui.py:3446,2236; testes_estados.json:alarm_frozen; estado_alarme_contador_congelado.png',
 'Modal aberto deve atualizar seu próprio evento ao menos na cadência do motor, inclusive após suspensão/retorno.')
finding('GUI-08','P1','Criticidade visual e métricas podem ficar desatualizadas',
 'DOCX §21: nome da faixa, cor e texto temporal coerentes.',
 'Evento a16min é Urgente; após avanço simulado de2min, contador muda mas badge continua Urgente, embora classificação calculada seja Crítico. Métricas dependem de nova renderização.',
 'Tick atualiza textos temporais; cores, classes, filtros temporais e indicadores só são recalculados por renderizações ocasionais.',
 'parts_js_ui.py:2236; parts_js_service.py:1268; testes_integridade_visual.json:band_stale',
 'Ao cruzar cada limite, atualizar classe, cor, nome, contagens e inclusão na lista urgente sem exigir interação ou alarme.')
finding('GUI-09','P1','Antecedência e volume zero não são respeitados',
 'Formulário admite antecedência0 e volume0%; comportamento deve corresponder ao valor exibido.',
 'Com antecedência global0, evento a10min disparou. Com volume0, Audio.play recebeu80. A tela permite uma configuração que não é executada.',
 'Uso de valor || padrão no salvamento e no motor transforma zero em15 ou80.',
 'parts_js_ui.py:1148; parts_js_service.py:1297,1367; testes_integridade_visual.json:zero_lead/zero_volume',
 'Distinguir ausência de zero; salvar/recarregar e disparar com0 sem substituição, mantendo checkbox de habilitação separado.')
finding('GUI-10','P0','Cancelar confirmação deixa a ação antiga armada',
 'Cancelar/fechar um diálogo deve descartar integralmente a ação em análise.',
 'AbrirA,cancelar,abrirB,confirmar resulta em ações[A,B]. Também há padrão semelhante em dependências/importação.',
 'Listener do botão confirmar só é removido na confirmação; fechar não remove closures anteriores.',
 'parts_js_ui.py:3598,3613,3629; testes_fluxos.json:confirmation_cancel',
 'Uma única ação atual por diálogo; cancelar,Escape e backdrop devem remover callbacks; confirmarB nunca executaA.')
finding('GUI-11','P1','Anotação salva não aparece no modal de detalhes',
 'DOCX §24 P0479 e §§6,15: desenho salvo visível nos detalhes.',
 'PNG existe; img display inline,container interno block,mas ancestral detailsModalDrawingBox permanece display:none. Imagem não tem retângulo visível.',
 'JS escolhe detailsStylusContainer e mostra somente o filho, sem exibir o pai oculto.',
 'parts_html.py:1843; parts_js_ui.py:3542; testes_fluxos.json:drawing_details; detalhes_com_desenho_oculto.png',
 'Um único container visível controlado de forma consistente; abrir eventos com e sem desenho alternadamente não deixa imagem residual ou oculta.')
finding('GUI-12','P1','Detalhe inicial de tablet/ultrawide abre sem conteúdo',
 'Designer.pdf e DOCX §15: seleção deve mostrar hierarquia, programação, descrição e ações.',
 'Inicialmente tab=details, painel mostra badge,contador,abas e botões,mas não os dados. Após clicar no evento,tab=detalhes e conteúdo aparece.',
 'Estado inicial usa chave inglesa details; renderPendingDetailPane compara detalhes,anotacoes,historico.',
 'parts_js_ui.py:13,2110,2127; testes_estados.json:detail_default/detail_clicked; estado_detalhe_tablet_inicial.png',
 'Normalizar chaves; seleção inicial e troca de modos devem manter aba/dados; teste sem clique adicional.')
finding('GUI-13','P1','Resumo do alarme mostra undefined min antes',
 'DOCX §15: detalhe selecionado apresenta programação e alarmes legíveis.',
 'Após corrigir a seleção por clique, o detalhe mostra Crítica • undefined min antes.',
 'Renderizador usa ev.leadTimeMinutes; modelo persiste alarmMinutesBefore e configurações globais podem definir outro valor efetivo.',
 'parts_js_ui.py:2144; parts_js_service.py:574; testes_estados.json:detail_clicked; estado_detalhe_tablet_apos_clique.png',
 'Mostrar antecedência efetiva correta para configuração padrão/específica; nenhum undefined/NaN na interface.')
finding('GUI-14','P1','Contador lateral de alarmes permanece zero',
 'DOCX §14 e referências: quantidade deve corresponder aos alarmes apresentados.',
 'Dois eventos triggered: cabeçalho2,métrica2,banner2,mas painel lateral0.',
 'HTML usa tabletAlarmsCount; JS procura badgeTabletAlarmsCount inexistente.',
 'parts_html.py:370; parts_js_ui.py:1709; testes_estados.json:alarm_badges; estado_alarmes_ativos_tablet.png',
 'Vincular o ID correto e definir se painel conta triggered ou também críticos/atrasados; título/lista/contagem coerentes.')
finding('GUI-15','P2','Cartões de concluídos não recolhem nem expandem',
 'DOCX §8: cartões compactos expansíveis no celular.',
 'Antes/depois do clique,detalhes continuam display:block e mesma altura; muda somente classe expanded.',
 'CSS estiliza completed-card-extra; HTML dinâmico usa completed-card-details.',
 'parts_css.py:1455; parts_js_ui.py:2371; testes_estados.json:completed_expansion; estado_concluidos_expansao_ineficaz.png',
 'Classe única; expansão por clique e teclado,aria-expanded e ação destrutiva protegida.')
finding('GUI-16','P0','Carregar demonstração esvazia a base e anuncia sucesso',
 'Botão deve carregar os exemplos anunciados e descrever claramente qualquer substituição.',
 'Clique real+confirmação:6eventos/4UOs passaram a0/0; toast diz Massa de dados de demonstração carregada com sucesso.',
 'Handler chama clearAllData,mas não chama loadDemoData antes de informar sucesso.',
 'parts_js_ui.py:1346; testes_estados.json:demo_before/demo_after; estado_botao_demo_base_esvaziada.png',
 'Após confirmar deve existir massa anunciada; se substituir dados,prévia e aviso devem dizer isso explicitamente,com recuperação verificável.')
finding('GUI-17','P0','Prévia de backup executa código fornecido pelo arquivo',
 'DOCX §§11,23: prévia de arquivo deve ser apresentação segura antes da importação.',
 'Um exportedAt inválido contendo um elemento com handler executou marcador JavaScript ao abrir o resumo,antes de confirmar.',
 'Formatador retorna texto inválido original e showImportSummaryModal insere o valor sem escape em innerHTML.',
 'parts_js_ui.py:3633; parts_js_core.py:111; testes_integridade_visual.json:backup_summary_script_execution',
 'Tratar texto como texto; validar esquema/formatos e URLs de imagem/IDs; arquivo inválido nunca cria atributos executáveis.')
finding('GUI-18','P0','Substituição de backup pode deixar a base vazia',
 'DOCX §§11,23 e preservação de integridade: validar antes de alterar e recuperar em falha.',
 'Backup com consoles:[{}] passa pela validação; replace limpa coleções e falha ao inserir chave ausente. No teste13consoles viraram0.',
 'Validação superficial; várias transações independentes após limpeza. Download preventivo não fornece rollback transacional.',
 'parts_js_service.py:1473,1504; testes_integridade_visual.json:replace_invalid (download preventivo substituído por stub apenas no teste)',
 'Validar todos registros/relações antes da confirmação final; substituição atômica; falha mantém base anterior íntegra.')
finding('GUI-19','P1','Mesclagem pode sobrescrever dado novo por antigo',
 'Promessa da interface/documentação: preservar registros mais recentes sem duplicar.',
 'Registro2026versão2 foi substituído por registro2020versão2.',
 'Merge compara apenas version com >=;updatedAt não participa e regras de negócio não são revalidadas.',
 'parts_js_service.py:1535; testes_integridade_visual.json:merge_old',
 'Política de conflito explícita e prévia dos conflitos; igualdade de versão não autoriza regressão silenciosa.')
finding('GUI-20','P1','CSV apresenta data e hora de conclusão inválidas',
 'DOCX §11 e tela Backup: formato brasileiro e arquivo utilizável.',
 'Download real contém 25T10:...Z/10/2026 e hora2026-. Entrada2026-10-25T15:30Z produz formatos inválidos. Strings iniciadas por= são mantidas como fórmulas potenciais.',
 'formatDate divide ISO por hífen;formatTime corta os cinco primeiros caracteres de qualquer string com dois-pontos;escapeCSV só protege delimitadores.',
 'parts_js_core.py:83,98,177; parts_js_service.py:1619; teste_previsoes_concluidas.csv; testes_integridade_visual.json:formats',
 'Tratar Date/ISO/data simples separadamente; fuso local consistente e política de texto para células de origem livre.')
finding('GUI-21','P1','Modais não recebem nem contêm o foco',
 'DOCX §19: teclado físico funcional,foco devolvido ao acionador e ações críticas protegidas.',
 'Ao abrir conclusão,foco permanece no botãoFiltros atrás do diálogo;Tab vai ao checkbox de console no fundo;Escape fecha,sem retornar ao acionador original.',
 'Abertura/fechamento só altera classes; não há controlador comum de foco,inert ou contenção deTab.',
 'parts_js_ui.py:3321,3460; testes_estados.json:focus_open/focus_tab/focus_escape',
 'Foco inicial no diálogo,fundo inerte,Tab circular,Escape controlado e retorno ao acionador; ensaio com leitor de tela.')
finding('GUI-22','P1','Menu fechado continua recebendo Tab fora da tela',
 'DOCX §§2.3,19,20: navegação acessível e controles alcançáveis.',
 'Sequência real deTab alcançou btnCloseMobileDrawer e todos os links do drawer fechado com coordenadas negativas.',
 'Menu fechado só recebe transform translateX; permanece exibido/focável; botão não anuncia expansão.',
 'parts_css.py:417; parts_js_ui.py:341; testes_estados.json:tab_sequence_drawer_closed',
 'Menu fechado deve sair da árvore de interação; aberto deve controlar foco e anunciar estado.')
finding('GUI-23','P2','Controles e diálogos sem nomes acessíveis suficientes',
 'DOCX §§5.1,19,20: campos e ícones com nome acessível e relação inequívoca.',
 'Inventário encontrou36controles sem label associado;2são auxiliares explicitamente ocultos. Nove dos dez diálogos não têm nome associado; nove fechamentos mostram só×.',
 'Labels adjacentes sem for; placeholders/símbolos substituem relações semânticas;aria-modal não nomeia diálogo.',
 'dados/inspecao_estatica.json; dados/analise_acessibilidade.md; parts_html.py:1701–1959',
 'Associar label/for ou aria-label/labelledby; nomes de ícones/fechamentos; validar leitura por tecnologia assistiva,sem supor que inventário sozinho certifica conformidade.',
 'Confirmado por inspeção estática; inventário completo disponível')
finding('GUI-24','P2','Campos,ações e textos ainda ficam menores que a orientação do DOCX',
 'DOCX §20 recomenda alvos~44px,campos48px,principal16px,auxiliar14px.',
 'Ordenação medida150×32px,fonte12px em celular/tablet; descrição tablet46px/14,4px; labels de métricas~11,5px; fechar modal36×36px. Descrição mobile48px/16px é positiva.',
 'Estilos inline pequenos prevalecem sobre regras responsivas e vários controles não usam classebtn.',
 'metricas_navegador.json:controls; parts_css.py:749,1080; parts_html.py:525,940; analise_acessibilidade.md',
 'Medir área realmente clicável,considerando labels decheckbox; adequar controles e hierarquia tipográfica sem encolher conteúdo essencial.')
finding('GUI-25','P2','Descrição permanece de uma linha e formulário não segue toda a hierarquia móvel',
 'DOCX §5.3 exige descrição multiline; §§5.1,16 organizam localização/programação/atividade/alarme em blocos.',
 'Descrição é INPUT type=text. Tablet não separa claramente localização e programação em duas regiões. No celular há barra de salvar,mas ainda excesso de ações e campos longos.',
 'Formulário reutiliza grade genérica de desktop,com descrição input e grupos extensos.',
 'parts_html.py:499; parts_css.py:593; testes_fluxos.json:description_type; celular_manual_viewNewEvent.png; tablet_manual_viewNewEvent.png',
 'Textarea,grupos lógicos coerentes,dois comandos de salvamento legíveis,campos preservados e política explícita paraEnter.')
finding('GUI-26','P2','Teclado virtual e áreas seguras têm cobertura incompleta',
 'DOCX §§19,22: barras não cobrem campo ativo; teclado apropriado; avançar não salva por acidente.',
 'Há tipos nativos date/time/number e safe-area na barra inferior,mas não há tratamento explícito de visualViewport,focusin,scrollIntoView,inputmode/enterkeyhint ou submissão implícita.',
 'Barras sticky/fixas dependem da geometria do viewport; apenas presença de CSS não demonstra comportamento com teclado.',
 'parts_html.py; parts_js_ui.py; analise_acessibilidade.md; screenshots nova previsão',
 'Testar Android/iOS com teclado aberto,retrato/paisagem/zoom; campo e ação principal visíveis. Não afirmar falha em todos aparelhos sem ensaio.',
 'Lacuna estática confirmada; efeito em teclado virtual não verificado')
finding('GUI-27','P2','Lousa preserva os bytes,mas sua proporção visual muda',
 'DOCX §§6,17: proporção preservada,prioridade stylus,ferramentas acessíveis,rotação sem perda.',
 'Texto e PNG foram preservados na rotação simulada. Porém canvasbitmap600×220 apareceu293×180 e depois747×180;mesmos pixels sofrem escalas distintas. Ferramentas têm positionstatic.',
 'Tamanho de buffer fixo e altura CSS independente da largura; fullscreen não recalcula buffer/proporção.',
 'parts_js_ui.py:410; parts_html.py:542; testes_fluxos.json:canvas_before/rotation/canvas_fullscreen; lousa_tela_cheia_paisagem.png',
 'Separar coordenadas lógicas e resolução de apresentação; preservar aspecto ou reamostrar sem deformar; verificar stylus,palma,pressão e nitidez em aparelho físico.')
finding('GUI-28','P2','Categoria de cadastro é perdida ao reaplicar visualização',
 'DOCX §1.2: mudança de layout preserva posição lógica e seleção.',
 'SelecionarTAGs e reaplicarTablet troca categoria paraConsoles.',
 'switchCadastrosCategory grava activeCadastrosCategory;applyResponsiveLayoutMode lê currentCadastrosCategory,que permanececonsoles.',
 'parts_js_ui.py:128,2435; testes_estados.json:category_before/category_after',
 'Estado único da categoria; alternar modo/rotação não troca contexto nem apaga edição em andamento.')
finding('GUI-29','P2','Atalhos das métricas não correspondem às listas abertas',
 'DOCX §3.4: cada métrica funciona como filtro/atalho coerente.',
 'Cards15/30/60min apenas navegam;Alarmes não reconhecidos aplica hasAlarm,que inclui programados/reconhecidos;Concluídos hoje não limita o dia;filtros anteriores podem continuar ativos.',
 'Handler só diferencia completed,overdue e unacknowledged;não define estado completo do filtro de destino.',
 'parts_js_ui.py:761; parts_html.py:270; analise_acessibilidade.md',
 'Lista aberta deve corresponder exatamente à contagem,com filtro explícito e sem restos de busca anterior;usar controles nativos acessíveis por teclado.',
 'Confirmado por inspeção do handler')
finding('GUI-30','P2','Hub de cadastros não entrega o fluxo hierárquico da referência',
 'DOCX §§9,18: celularcategoria→lista→formulário;tabletduas regiões.',
 'Categorias filtram seções,mas lista e formulário continuam no mesmo hub vertical. Botões de categoria aparecem sem a aparência definida para cad-cat-btn.',
 'HTML cadastros-category-btn diverge de CSS cad-cat-btn;gridTablet é uma coluna e categorias não formam rail lateral.',
 'parts_html.py:895; parts_css.py:1474,1849; celular_manual_viewCadastros.png; tablet_manual_viewCadastros.png',
 'Coerência entre classes e apresentação; categoria/lista/formulário claros e estado de edição preservado.')
finding('GUI-31','P2','Desktop convencional não possui navegação lateral fixa',
 'DOCX §1.2 P0023 e Designer(2).pdf: navegação lateral,indicadores,destaque e lista.',
 'ModoConvencional mantémapp-navhorizontal; não existe shell lateral persistente. Há boa tabela e métricas,mas composição diverge.',
 'Não há implementação de sidebar desktop;modo apenas mantém apresentação base.',
 'parts_html.py:209; modo_conventional_viewDashboard.png; Designer(2).pdf página1',
 'Definir shell convencional conforme referência ou registrar formalmente exceção; não avaliar somente pela existência do seletor.')
finding('GUI-32','P2','Widescreen e ultrawide ampliam o espaço sem compor todas as regiões',
 'Designer(2).pdf e DOCX §§1.1–1.2: lista,detalhe,métricas/contexto e ações em regiões coordenadas.',
 'Wide mostra destaque+alarmes e tabela abaixo. Ultra estica a composição;Pendentes tem lista/detalhe,mas não reúnetodas regiões. Gráficos e rail de ações da prancha não existem.',
 'CSS amplia max-width e algumas grids; não há componentes para todos os painéis da referência.',
 'parts_css.py:1869; modo_widescreen_viewDashboard.png; modo_ultrawide_viewDashboard.png; matriz_visual.json',
 'Aproveitar área para contexto coordenado e listas legíveis,evitando colunas de texto excessivamente largas. Gráficos são divergência de referência,não requisito operacional autônomo inferido.')
finding('GUI-33','P2','Temas existem,mas elementos críticos ainda escapam dos tokens',
 'Designer(1).pdf: identidade coerente dos três temas,com prioridade à leitura da criticidade.',
 'Após estabilizar transições,Industrial tem botãoamarelo/textopreto e navegação clara,corretamente. Persistem azul fixo doSpotlight/header/abaativa e contadoratrasado#991b1b em superfícies escuras. Claro mantém gradiente destacado e texto auxiliar pequeno.',
 'Muitos estilos inline/hexadecimais em lugar dos tokens;classe de evento também fornece cor sem adequação ao fundo.',
 'parts_css.py:115,387; parts_js_ui.py:1756; testes_estados.json:theme_colors; tema_industrial_viewDashboard.png; tema_light_viewDashboard.png',
 'Revisar texto/fundo real por componente e estado;medir contraste e validar luminosidade física. Não declarar tema inteiro inválido nem copiar literalmente cores de imagem gerada.')
finding('GUI-34','P2','Configurações e cores não seguem completamente a simplificação móvel',
 'DOCX §§10–11: grupos expansíveis e edição de faixas adequada ao celular.',
 'Configurações de alertas permanecem longa página;faixas têmcards,mas editor continua na mesma visão e cadeia só reflete dados já salvos. Highlight/pulsing são salvos,mas não governam os renderizadores inspecionados.',
 'Controles presentes sem vinculação integral à apresentação;ausência de composição expansível de todos os grupos.',
 'parts_js_ui.py:1153,3108; partes HTML/CSS; celular_manual_viewConfigAlerts.png; celular_manual_viewConfigColors.png; analise_docx.md',
 'Grupos claros,preview de edição e opção visual realmente respeitada;movimento reduzido deve continuar funcionando independentemente da preferência salva.')
finding('GUI-35','P2','Listas de filtros podem manter cadastros antigos',
 'Seletores devem refletir criação,edição,inativação e importação de cadastros.',
 'populateFilterDropdowns só preenche se options.length<=1. Após primeiro preenchimento,novos registros/nomes podem não aparecer até recarregar.',
 'Cache implícito baseado no número de opções,sem invalidação.',
 'parts_js_ui.py:2256; analise estática dos seletores',
 'Repopular com preservação da seleção válida após cada alteração/importação;não deixar itens removidos nem impedir selecionar novos registros.',
 'Confirmado por inspeção da condição; cenário completo de alteração requer regressão')
finding('GUI-36','P2','Diagnóstico e avisos não garantem refletir toda a situação real',
 'DOCX §§12,23: status compreensível,limitações do navegador e dados locais explícitos.',
 'Há statusmodal,avisobackup e documentação offline,positivos. Diagnóstico segue tabela densa;relatório declara v1.0.0 versus UIv1.1.0;logError não é usado nos fluxos inspecionados. Lembrete usa7dias fixos apesar da chave configurável.',
 'Identificação de versão duplicada;tratamento de erro no console sem alimentar coleção;prazo hardcoded.',
 'parts_js_ui.py:228,3296; parts_js_core.py:648; celular_manual_viewDiagnostic.png; analise_docx.md',
 'Versão centralizada,erros úteis persistidos,lista móvel legível e período efetivo coerente com configuração;avisos não prometem garantia de alarme em segundo plano.')
finding('GUI-37','P2','Ações contextuais e identificação do controlador ficam incompletas',
 'DOCX §§3.3,4.2,8: contexto claro e histórico confiável.',
 'Folha de ações deixa subtítulo--. Concluídos usa completedBy||createdBy||Usuário local,mas conclusão não grava esses campos,portanto pode apresentar atribuição genérica em lugar do operador do histórico.',
 'Markup auxiliar não é preenchido;renderizador espera propriedades que o modelo não persiste.',
 'parts_html.py:1950; parts_js_ui.py:2205,2353; parts_js_service.py:622; celular_manual_modal_acoes.png',
 'Contextualizar evento no sheet;mostrar responsável real ou declarar ausência,nunca inferir identidade a partir de fallback genérico.')
finding('GUI-38','P2','Referências documentais têm ambiguidades internas',
 'Uma especificação coerente deve definir composição e prioridade sem leituras conflitantes.',
 'Tabela inicial:desktop>1024px com navegaçãohorizontal;§1.4:tablet até1199px;§1.2 e PDF:desktoplateral. Pranchaultrawide tem resolução ilustrativa inconsistente com razão32:9.',
 'Versões incrementais do documento e rótulos de mockup não foram harmonizados.',
 'DOCX tabela0/P0023/P0032; Designer(2).pdf página1; inventario_referencias.json',
 'Consolidar critério por conteúdo e composição. Não tratar a diferença isolada1600vs1920como falha mandatória:breakpoints são explicitamente recomendações.',
 'Confirmado nas referências; decisão de produto pendente')
finding('GUI-39','P3','Lista grande renderiza cartões e tabela simultaneamente',
 'Continuidade de operação e navegação legível com histórico/pendências crescentes.',
 'Teste sintético503eventos:503cards e503linhas criados,embora um modo fique oculto;tempo de renderização registrado no JSON. Não há paginação/virtualização.',
 'renderPendingEvents monta integralmente duas apresentações com getAll.',
 'parts_js_ui.py:1876,2057; testes_estados.json:render_500',
 'Definir volume esperado;medir em tablet real;renderizar apresentação necessária e limitar lotes se houver degradação. Tempo headless desta máquina não é benchmark de aparelhos.')
finding('GUI-40','P2','Fechamento genérico e ações destrutivas exigem revisão de interação',
 'DOCX §§19–20: evitar fechamento acidental e separar ações destrutivas das frequentes.',
 'Backdrop fecha qualquermodal;Escape remove todos os overlays ativos. Excluir continua ao lado de ações comuns em várioscards/linhas. Adiar5m abre seletor em vez de aplicar5min;input1–1440 não passa por checkValidity.',
 'Controlador de modal genérico sem política poração;ações replicadas em templates distintos.',
 'parts_js_ui.py:3321,3356,1729,2048; analise_acessibilidade.md; celular_manual_modal_adiamento.png',
 'Política explícita por diálogo,proteção durante escrita,ações destrutivas separadas e limites de adiamento validados;texto do comando deve descrever seu efeito.')

SCREENS=[
 ('Painel principal','Parcial com falhas críticas de coerência temporal','Métricas,Spotlight e urgentes existem; no Auto móvel a ordem segue desktop. Há contador de outro evento,classificação que pode ficar antiga e faixa de alarme sobreposta. Em tablet contador lateral0; em wide/ultra faltam regiões da referência.'),
 ('Nova previsão','Parcial','Seletores encadeados,campos nativos,ações salvar e desenho existem. Descrição de linha única,campos/tablet pequenos,formulário sem toda a separação de regiões e dependência de barras fixas ainda não validada com teclado físico móvel.'),
 ('Eventos pendentes','Parcial','Pesquisa,ordenação,cartões,tabela e detalhes presentes. Filtros não viram sheet,excesso de ações e detalhe inicial vazio. Tabela preferida pode continuar no celular;há risco de filtros desatualizados.'),
 ('Eventos concluídos','Parcial','Tabela de novecolunas foi corrigida e cards existem. Expansão ineficaz,identificação de controlador genérica e CSV de conclusão inválido;atalhoHoje não limita o período.'),
 ('Cadastros operacionais','Parcial','Cinco tipos de cadastro,filtros ativo/inativo e criaçãoinline presentes. Fluxo móvel/tablet segue hub e não lista→formulário/duasregiões;classes dos botões divergem;categoria pode ser perdida ao mudar modo.'),
 ('Configurações de alertas','Parcial','Controles amplos e três temas existem,mas página não é organizada em todas as seções expansíveis;zeros mudam de valor e algumas preferências visuais não comandam renderização.'),
 ('Configurações de cores','Parcial','Cards e cadeia existem,com validação de faixas. Edição não tem toda a prévia/organização prevista e mudanças temporais não garantem atualização dos cartões já exibidos.'),
 ('Backup e dados','Não pronta para validação operacional','Downloads funcionam noEdge,mas prévia executa código,replace pode esvaziar base,merge pode regredir,demonstrativo apaga dados com sucesso enganoso e CSV tem dadosinválidos. Importação móvel segue modalgenérico.'),
 ('Diagnóstico técnico','Parcial','Contagens e cópia existem. Tabela permanece densa,versão divergente e ausência de uso efetivo do registrodeerros. Status não comprova alertas sob bloqueio/suspensão.')
]

def human(s):
    if not isinstance(s,str):return s
    s=re.sub(r'(?<=\d)(?=[A-Za-zÀ-ÿ])',' ',s)
    s=re.sub(r'(?<=[a-uw-zà-ÿ])(?=\d)',' ',s)
    s=re.sub(r',(?=[^\s\d])',', ',s)
    s=re.sub(r';(?=\S)','; ',s)
    for a,b in {'CelularAuto':'Celular Auto','CelularManual':'Celular Manual',
                'TabletAuto':'Tablet Auto','TabletManual':'Tablet Manual','DesktopAuto':'Desktop Auto',
                'TemaProfissionalClaro':'Tema Profissional Claro','TemaIndustrial':'Tema Industrial',
                'doSpotlight':'do Spotlight','daTab':'da Tab','doCSV':'do CSV','deTab':'de Tab',
                'deJSON':'de JSON','de2':'de 2','de7':'de 7','de10':'de 10',
                '≤1':'≤ 1','conjuntode':'conjunto de','dadose':'dados e',
                'textoamarelo':'texto amarelo','botãoamarelo':'botão amarelo','textopreto':'texto preto',
                'superfíciesescuras':'superfícies escuras','todosregistros':'todos os registros',
                'alarmesativos':'alarmes ativos','tabletduas':'tablet duas','celularcategoria':'celular categoria',
                'campos/tablet':'campos no tablet','textoTablet':'texto Tablet',
                'campoFiltros':'campo Filtros','botãoFiltros':'botão Filtros',
                'concluídoshoje':'concluídos hoje','tempoheadless':'tempo headless',
                'header/nav/sheet':'header / nav / sheet','semestre':'semestre'}.items():s=s.replace(a,b)
    for a,b in {'Designer(1)':'Designer (1)','Designer(2)':'Designer (2)',
                'eDesigner':'e Designer','eDesktop':'e Desktop','e(2)':'e (2)',
                'novecolunas':'nove colunas','atalhoHoje':'atalho Hoje','criaçãoinline':'criação inline',
                'duasregiões':'duas regiões','noEdge':'no Edge','dadosinválidos':'dados inválidos',
                'modalgenérico':'modal genérico','registrodeerros':'registro de erros',
                'descriçãoTablet':'descrição no tablet','eventostriggered':'eventos triggered',
                'alarmestriggered':'alarmes triggered','diálogos:alarme':'diálogos: alarme',
                'relatório,as':'relatório, as','dados; scripts':'dados; os scripts',
                'concluídos,apesar':'concluídos, apesar','Semáudio':'Sem áudio',
                'semTab':'sem Tab','semEnter':'sem Enter','softwaremóvel':'software móvel',
                'páginafull':'página full','deúltimo':'de último','apenasA':'apenas A',
                'confirmarB':'confirmar B','executaA':'executa A','abrirA':'abrir A','abrirB':'abrir B',
                'modoAuto':'modo Auto','painelTablet':'painel Tablet','naAuto':'no Auto',
                'em19':'em 19','doiscomandos':'dois comandos','geraisde':'gerais de',
                'masa':'masa','novevisões':'nove visões','camposnativos':'campos nativos',
                'CSVpendentes':'CSV pendentes','CSVconcluídos':'CSV concluídos','CSVcadastros':'CSV cadastros',
                'mesmoHTML':'mesmo HTML','eIndexedDB':'e IndexedDB','comaria-live':'com aria-live',
                'peloEdge':'pelo Edge','apósreload':'após reload','alvosP1':'alvos P1',
                'telaútil':'tela útil','seção24':'seção 24','matrizDOCX':'matriz DOCX',
                'matrizvisual':'matriz visual','tipoP0':'tipo P0','objetosCSS':'objetos CSS',
                'dostrês':'dos três','Estático':'Estático'}.items():s=s.replace(a,b)
    return s

def csv_export(name,rows,fields):
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields,delimiter=';',extrasaction='ignore');w.writeheader()
        for row in rows:
            w.writerow({k:' | '.join(map(str,v)) if isinstance(v,list) else v for k,v in row.items() if k in fields})

def main():
    for f in F:
        for key in ['titulo','esperado','observado','causa','aceite']:
            f[key]=human(f[key])
    req=read('matriz_requisitos_docx.json');visual=read('matriz_visual.json')
    # Final cross-review: optional alternatives are not cumulative obligations;
    # absence of an explicit handler does not prove failure on physical keyboards.
    for x in req:
        origin=x['origem']
        if origin in ['P0398','P0441','P0442']:
            x['status']='nao_verificado'
            x['observacao']='Não há tratamento explícito de visualViewport/rolagem ao foco no código. O navegador pode aplicar comportamento nativo; cobertura real de campos por teclado virtual e barras fixas exige aparelho físico e não foi comprovada nesta auditoria.'
        elif origin=='P0445':
            x['status']='parcial'
            x['observacao']='Pesquisa reage ao evento input, portanto digitação pelo teclado já filtra; não há enterkeyhint/ação móvel explícita. Ausência desse atributo não significa busca inoperante. Interação com teclado virtual físico permanece pendente.'
        elif origin=='P0463':
            x['status']='parcial'
            x['observacao']='Downloads JSON e três CSVs exercitados pela GUI no Edge em perfil temporário; importações inválidas e conflito foram reproduzidos em testes isolados. Android/iOS e navegadores-alvo reais continuam pendentes. Ver testes_estados.json e testes_integridade_visual.json.'
        elif origin=='P0156':
            x['status']='parcial'
            x['evidencia_codigo']=['parts_html.py:746','parts_js_ui.py:1536','parts_css.py:1727']
            x['observacao']='Botões Limpar/Aplicar existem e têm handlers. O conjunto é parcial porque o painel não ganha a classe active-sheet e permanece no fluxo; não se afirma ausência dos botões.'
        elif origin in ['P0161','P0162']:
            x['status']='parcial';x['natureza']='recomendacao'
            x['observacao']='Alternativas não cumulativas: tabela existe como opção com container rolável, porém preferência de tabela pode continuar no celular sem apresentação explícita como modo avançado. Não se exige ocultar e oferecer tabela simultaneamente; o problema principal do padrão móvel está em P0160.'
        elif origin in ['P0316','P0318']:
            x['status']='aderente'
            x['observacao']='Alternativas não cumulativas para tablet em retrato. Drawer retrátil é a alternativa implementada e atende a possibilidade do conjunto; não se exige também barra estreita. O Auto ainda não ativa a composição tablet, conforme GUI-01; sidebar persistente paisagem é outro requisito.'
        elif origin=='P0281':
            x['status']='parcial'
            x['observacao']='Form-grid expande controles pela coluna e slider está presente; há grupos de repetição lado a lado e áreas de checkbox próprias. A geometria varia por perfil (metricas_navegador.json). A ausência de seções expansíveis não é evidência para reprovar a largura dos controles.'
        elif origin=='P0282':
            x['status']='parcial';x['evidencia_codigo']=['parts_html.py:1255','parts_html.py:1324']
            x['observacao']='Timbre e volume ficam antes de canais/repetição; Testar Som fica na barra de ações ao final, separado por esses grupos. Testar Notificação fica próximo ao botão de permissão. Avaliação parcial pela distribuição efetiva, não pela ausência de accordion.'
    write('matriz_requisitos_docx.json',req)
    # Keep all 90 screenshots in a single index, including supplementary states.
    index=read('capturas_indice.json');indexed={x['file'] for x in index}
    from PIL import Image
    for path in sorted((OUT/'capturas').glob('*.png')):
        if path.name not in indexed:
            with Image.open(path) as im:w,h=im.size
            index.append({'file':path.name,'profile':'estado_complementar','view':path.stem,
                          'width':None,'height':None,'screenshot_width':w,'screenshot_height':h})
    write('capturas_indice.json',index)
    metrics=read('metricas_navegador.json');flows=read('testes_fluxos.json');states=read('testes_estados.json')
    stat=read('inspecao_estatica.json');integrity=read('testes_integridade_visual.json')
    write('achados_consolidados.json',F)
    csv_export('MATRIZ_REQUISITOS.csv',req,['id','secao','natureza','requisito','status','evidencia_codigo','observacao','prioridade'])
    csv_export('MATRIZ_VISUAL.csv',visual,['id','referencia','elemento','expectativa','status','evidencia','observacao'])
    csv_export('ACHADOS_PRIORIZADOS.csv',F,['id','prioridade','titulo','tipo','esperado','observado','causa','evidencia','aceite'])
    fingerprints=[]
    for path in [ROOT/'index.html',*sorted(ROOT.glob('parts_*.py')),ROOT/'build_all.py',*sorted(ROOT.glob('*.docx')),*sorted(ROOT.glob('Designer*.pdf'))]:
        fingerprints.append({'file':path.name,'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    write('manifesto_fontes.json',fingerprints)
    count_status=Counter(x['status'] for x in req);count_nature=Counter(x['natureza'] for x in req)
    docx_notes=(DATA/'analise_docx.md').read_text(encoding='utf-8')
    docx_notes=re.sub(r'A matriz registra \d+ itens:.*?\.',
        'A matriz registra '+str(len(req))+' itens: '+', '.join(k+': '+str(v) for k,v in count_status.items())+'.',docx_notes,count=1)
    if 'Revisão final integrada' not in docx_notes:
        docx_notes+='\n\n## Revisão final integrada\n\nA consolidação posterior limitou afirmações sobre teclado virtual não testado, reconheceu busca por input e botões Limpar/Aplicar existentes, tratou alternativas de navegação/tabela como não cumulativas e incorporou os testes de downloads no Edge. A matriz JSON e o relatório consolidado contêm os estados finais; os testes físicos permanecem pendentes.\n'
    (DATA/'analise_docx.md').write_text(docx_notes,encoding='utf-8')
    widths=sorted({x['width'] for x in metrics});shots=sorted((OUT/'capturas').glob('*.png'))
    summary={'date':'2026-10-25','timezone':'America/Sao_Paulo','browser':flows['browser'],'findings':len(F),
             'priority':dict(Counter(x['prioridade'] for x in F)),'requirements':len(req),'requirements_status_static':dict(count_status),
             'requirements_nature':dict(count_nature),'visual_items':len(visual),'measurements':len(metrics),'widths':widths,'screenshots':len(shots),
             'scope_note':'Matriz DOCX é avaliação estática de itens rastreáveis, incluindo recomendações e sínteses repetidas. Não é percentual de aprovação funcional.'}
    write('resumo_auditoria.json',summary)
    blocks=[]
    def heading(t):blocks.append(('h',human(t)))
    def para(t):blocks.append(('p',human(t)))
    def table(headers,rows):blocks.append(('table',(headers,[[human(c) for c in row] for row in rows])))
    def figure(file,caption):blocks.append(('figure',(file,human(caption))))
    heading('Auditoria completa da interface gráfica')
    para('Caderno de Previsões Operacionais em Dutos | 25/10/2026 | Horário de Brasília | Revisão documental,visual,estática e de execução. Fontes da aplicação preservados; dados de teste pertencem a perfis temporários do navegador.')
    heading('1. Parecer')
    para('A interface é parcialmente coerente com o documento e com as referências. Há uma base funcional reconhecível,mas a adaptação móvel/tablet não está concluída e a composição de monitores amplos não corresponde integralmente às pranchas. Há falhas de apresentação que alteram o significado operacional do que é exibido; portanto não se trata apenas de refinamento estético. Não recomendo considerar esta versão aprovada contra a especificação antes de corrigir os bloqueadores e repetir os critérios de aceite deste relatório.')
    para(f'O relatório registra 40 grupos de achados, 321 itens rastreáveis do DOCX, 22 elementos de comparação visual, 90 capturas e 198 medições de tela em 19 larguras. As contagens dos requisitos não equivalem a 321 obrigações únicas: {count_nature["obrigacao"]} itens foram classificados como obrigações, {count_nature["recomendacao"]} como recomendações e {count_nature["sintese"]} como sínteses que repetem requisitos anteriores. Não foi atribuído percentual subjetivo de fidelidade.')
    table(['Prioridade','Interpretação','Grupos'],[['P0','Ação inesperada,perda de dados ou execução de conteúdo pelo arquivo',summary['priority'].get('P0',0)],['P1','Informação operacional incoerente ou ação essencial inacessível',summary['priority'].get('P1',0)],['P2','Composição,acessibilidade e fluxo parcialmente implementados',summary['priority'].get('P2',0)],['P3','Melhoria de escala e manutenção da experiência',summary['priority'].get('P3',0)]])
    heading('2. Materiais e método')
    para('O DOCX contém520parágrafos e1tabela. Foram extraídos478parágrafos não vazios,com marcadores P0000…P0519,além de6linhas da tabela. A cobertura preserva parágrafos normativos,exemplos,contexto e justificativas de inclusão/exclusão. O arquivo cobertura_paragrafos_docx.json permite localizar cada trecho; requisitos_extraidos.txt mantém a transcrição textual. Os três PDFs têm uma página cada e foram renderizados e examinados visualmente,sem presumir equivalência pelos nomes dos PNGs.')
    table(['Arquivo','Conteúdo avaliado'],[['Designer.pdf','Seis fluxos de celular e dois de tablet: painel,pendentes,novo,alarme,filtros,drawer e lista/detalhe.'],['Designer (1).pdf','Controle Escuro,Profissional Claro e Industrial Alto Contraste; exemplos de painel,listas e formulário.'],['Designer (2).pdf','Monitor convencional,widescreen e ultrawide;regiões coordenadas,navegação lateral,contexto e ações.'],['Imagens adicionais','Designer (1).png e(2).png são duplicatas entre si;Designer (3).png eDesigner.jpg mostram temas;Designer.jpeg é inteiramente branco. PDF e PNG com mesmo sufixo não são intercambiáveis.']])
    para('Wireframes são referências de comportamento,conforme o próprio DOCX. Foram cobrados hierarquia,composição,prioridade operacional e ergonomia; não foram cobrados pixel a pixel,carcaças de aparelhos,tipografia distorcida das imagens ou textos ilustrativos gerados. Os gráficos da prancha são divergências de referência visual,sem serem promovidos automaticamente a novas regras operacionais.')
    para(f'Execução: Microsoft Edge/Chromium {flows["browser"]},Playwright,arquivo local file://,perfil temporário e timezone America/Sao_Paulo. Nove visões foram abertas nos perfis CelularAuto390×844,CelularManual390×844,TabletAuto820×1180,TabletManual1024×768 eDesktopAuto1440×1000. Foram avaliadas ainda as larguras '+', '.join(map(str,widths))+'px; os sete valores de visualização foram persistidos e recarregados; os três temas foram capturados após estabilização das transições.')
    para('Para comparação determinística,o motor periódico foi pausado após um tick inicial; eventos de demonstração e alarmes foram criados apenas nas bases temporárias. Testes temporais avançam Date.now sem alterar o relógio do sistema. Não se deve interpretar relógio parado na captura como falha da aplicação. Capturas finais desativam animações durante a imagem para evitar estados transitórios de cor. Imagens full-page documentam página inteira; elementos sticky/fixed aparecem na posição de captura e devem ser lidos junto das medidas. A imagem de rolagem do banner é de viewport,para evidenciar a sobreposição real.')
    heading('3. Resultado por tela')
    table(['Visão','Avaliação','Resultado'],SCREENS)
    heading('4. Evidências de composição e responsividade')
    profiles=[]
    for label,title in [('celular_auto','Celular automático'),('celular_manual','Celular manual'),('tablet_auto','Tablet automático'),('tablet_manual','Tablet manual'),('desktop_auto','Desktop automático')]:
        rows=[x for x in metrics if x['profile']==label]
        profiles.append([title,str(rows[0]['width'])+'px',str(min(x['scrollWidth'] for x in rows))+'–'+str(max(x['scrollWidth'] for x in rows))+'px','Extrapola página inteira' if any(x['scrollWidth']>x['width']+2 for x in rows) else 'Sem extrapolação global neste conjunto'])
    table(['Perfil','Viewport','Largura do conteúdo nas nove telas','Resultado'],profiles)
    figure('celular_auto_viewDashboard.png','CelularAuto390px: métricas antes do destaque,navegação desktop e conteúdo1077px.')
    figure('celular_manual_viewDashboard.png','CelularManual390px: ordem melhora,mas cabeçalho extrapola e barra inferior não recebe a aparência prevista.')
    figure('tablet_manual_viewPendingEvents.png','TabletManual1024px: lista/detalhe existe,porém estado inicial de aba pode deixar o painel sem dados.')
    figure('modo_ultrawide_viewDashboard.png','Ultrawide3440px: a composição estica,sendo diferente das regiões coordenadas mostradas em Designer(2).pdf.')
    heading('5. Temas,ergonomia e acessibilidade')
    para('Os três temas estão implementados e são distinguíveis. O Industrial aplica fundo preto, superfícies escuras, texto claro e botão primário amarelo/preto. Persistem cores fixas azuis e vermelho escuro no contador. A aprovação deve medir os pares reais de texto/fundo e ser complementada por ensaio de leitura e legibilidade em aparelhos físicos.')
    para('A classificação utiliza texto e ícone além de cor,aspecto positivo. Emulação prefers-reduced-motion produziu animationDuration0s nos cards e badges pulsantes. Por outro lado,a ordenação tem32px de altura e12px de fonte;descriçãoTablet46px/14,4px;há labels de métricas e auxiliares abaixo dos14px recomendados. O inventário de controles informa dimensões reais por perfil;checkboxes devem ser avaliados com sua label e área acionável,sem contar apenas o quadrado.')
    para('A inspeção encontra34controles visíveis ou condicionais sem associação de label,além de2auxiliares ocultos;9de10diálogos sem nome associado e9fechamentos sem nome explícito. Isso é inventário estático,com placeholders/title registrados separadamente,não certificação formal de acessibilidade. Testes deTab confirmam foco atrás de modal e em drawer fechado fora da tela. Leitor de tela e teclado virtual físico permanecem pendentes.')
    figure('tema_light_viewDashboard.png','TemaProfissionalClaro com transições estabilizadas;comparar com coluna central de Designer(1).pdf.')
    figure('tema_industrial_viewDashboard.png','TemaIndustrial estabilizado: botão principal amarelo confirmado;Spotlight e marcações azuis continuam presentes.')
    heading('6. Achados detalhados e critérios de aceite')
    for f in F:
        heading(f'{f["id"]} | {f["prioridade"]} | {f["titulo"]}')
        para('Base: '+f['esperado'])
        para('Resultado: '+f['observado'])
        para('Causa identificada: '+f['causa'])
        para('Evidência: '+f['evidencia']+' | Classificação: '+f['tipo'])
        para('Critério de aceite: '+f['aceite'])
    heading('7. Modais,estados e ações conferidos')
    para('Foram abertos os dez diálogos:alarme,adiamento,conclusão,reabertura,detalhes,confirmação,dependências,importação,ações contextuais eStatus. Também foram capturadosdrawer,filtros,lousa em tela cheia,alarmestriggered,estado vazio,texto longo,desenho salvo invisível e detalhe antes/depois de selecionar. Presença e abertura não significam aprovação:os achados especificam foco,contador,fechamento e geometria.')
    figure('celular_manual_modal_filtros.png','Filtro móvel:posiçãoestática,sem a classeactive-sheet que ativa o painel inferior.')
    figure('estado_banner_rolagem.png','Viewport após rolagem:cabeçalho de altura variável cobre a faixa posicionada emtop52px.')
    figure('estado_detalhe_tablet_apos_clique.png','Após selecionar,o detalhe aparece,mas antecedência mostra undefined min antes.')
    figure('detalhes_com_desenho_oculto.png','Desenho salvo não fica visível no modal porque o ancestral continua oculto.')
    heading('8. As doze correções obrigatórias da seção24')
    fixes=[x for x in req if x['origem'] in {f'P{i:04d}' for i in range(470,482)}]
    table(['Origem','Correção','Situação','Evidência'],[[x['origem'],x['requisito'],x['status'],x['observacao']] for x in fixes])
    para('Onze correções têm equivalência funcional no código;uma não está resolvida:desenho de detalhes. Essa contagem é restrita às12correções preexistentes,não à prontidão da interface. Definir aliasCSS é solução funcional válida,assim como normalizar concluido→completed no serviço. A regra de um únicoNovo está presente,mas o estilo da barra inferior continua defeituoso.')
    heading('9. Resultados positivos e limites')
    para('Confirmados:mesmoHTML eIndexedDB;372IDs únicos;IDs do fonte eindex coincidentes;seletores encadeados comaria-live;três temas;sete modos persistidos apósreload;quatro downloads reais peloEdge(JSON,CSVpendentes,CSVconcluídos,CSVcadastros);movimento reduzido;preservação de texto ePNG na rotação simulada;abertura dasnovevisões sem erroJavaScript não tratado no ensaio. O conteúdo doCSVconcluídos,apesar de baixar,está incorreto.')
    para('O ensaio sintético de503pendentes criou503cards e503linhas de tabela;tempo observado nesta execução:'+str(states['render_500']['ms'])+'ms. Isso demonstra duplicação de apresentação,sem provar desempenho aceitável em tablet. Não se mediramRAM,CPU contínua,consumo de bateria ou funcionamento prolongado.')
    para('Pendências obrigatórias de homologação:AndroidChrome,iOSSafari,dispositivo com caneta e palma reais,pressão e precisão,luz solar,leitor de tela,zoom de sistema/200%,teclado virtual,áreas seguras de aparelhos,navegador em segundo plano,tela bloqueada,suspensão do computador,permissões reais de notificação,saída sonora física e múltiplas abas. O ensaio não afirma compatibilidade integral com esses cenários. Emular viewport não emula o sistema operacional móvel.')
    heading('10. Ordem de correção e regressão')
    table(['Etapa','Trabalho','Condição para avançar'],[['1','P0:cancelamento,prévia/importação,demonstração','Nenhuma ação cancelada executa;nenhum arquivo executa código;falha não altera base anterior.'],['2','P1:contador único,modal,classificação,zeros,desenho,detalhes,badge','Descrição/tempo/estado correspondem ao mesmo evento;UI reflete o banco e as preferências.'],['3','Shell móvel/tablet e filtros','Sem overflow global;header/nav/sheet/ações críticas visíveis em19larguras e nas rotações.'],['4','Acessibilidade e formulários','Fluxo teclado completo;nomes acessíveis;alvos/fontes conforme orientação;ensaio com teclado virtual.'],['5','Desktop amplo e acabamento dos temas','Composição validada contra cadaPDF;cores medidas;regiões necessárias presentes.'],['6','Homologação física e volume esperado','Aceites executados em aparelhos-alvo,com evidência por browser,orientação e estado.']])
    heading('11. Como consultar os registros')
    para('RELATORIO_INTERFACE.html reúne o texto integral,a matrizDOCX completa,a matrizvisual e a galeria das90capturas com busca/filtro locais. RELATORIO_INTERFACE.docx inclui o relatório,figuras selecionadas e as duas matrizes completas. RELATORIO_INTERFACE.md mantém versão textual. MATRIZ_REQUISITOS.csv,MATRIZ_VISUAL.csv eACHADOS_PRIORIZADOS.csv permitem filtrar e priorizar em planilha. A pasta dados contém JSONs de medições,fluxos,estados,inventário ehashes; scripts reproduz a auditoria sem editar a aplicação. Downloads de teste são massa fictícia criada no perfil temporário.')
    para('As classificações da matrizDOCX são predominantemente estáticas;os achados GUI-01…GUI-40 e os JSONs de testes fornecem validação de execução e prevalecem em caso de diferença. Um requisito pode ter estrutura presente e falhar noAuto ou em outra parte do fluxo;por isso não se deve usar uma linhaaderente como aprovação transversal da funcionalidade.')
    # Markdown output.
    md=[]
    for typ,value in blocks:
        if typ=='h':md.append('## '+value+'\n')
        elif typ=='p':md.append(value+'\n')
        elif typ=='figure':md.append(f'![{value[1]}](capturas/{value[0]})\n')
        else:
            headers,rows=value
            md.append('| '+' | '.join(headers)+' |\n| '+' | '.join('---' for _ in headers)+' |')
            md.extend('| '+' | '.join(str(c).replace('|','/').replace('\n',' ') for c in row)+' |' for row in rows)
            md.append('')
    md.append('## Apêndice A — Matriz integral do DOCX\n')
    md.append('| ID / natureza | Requisito | Situação / evidência / observação |\n| --- | --- | --- |')
    for x in req:md.append('| '+x['id']+' / '+x['natureza']+' | '+x['requisito'].replace('|','/')+' | '+x['status']+'; '+', '.join(x['evidencia_codigo'])+'; '+x['observacao'].replace('|','/')+' |')
    (OUT/'RELATORIO_INTERFACE.md').write_text('\n'.join(md),encoding='utf-8')
    # Word, with full requirements and reference matrices.
    doc=Document();sec=doc.sections[0];sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(6)
    footer=sec.footer.paragraphs[0];footer.text='Auditoria da interface | 25/10/2026 | '
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for typ,value in blocks:
        if typ=='h':doc.add_heading(value,1 if re.match(r'^\d+\.',value) else 2)
        elif typ=='p':doc.add_paragraph(value)
        elif typ=='figure':
            from PIL import Image
            path=OUT/'capturas'/value[0]
            with Image.open(path) as im:w,h=im.size
            target_w=min(6.5,7.0*w/h)
            doc.add_picture(str(path),width=Inches(target_w));doc.add_paragraph(value[1],'Caption')
        else:
            headers,rows=value;t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
            for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
            for row in rows:
                for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    doc.add_page_break();doc.add_heading('Apêndice A — Matriz integral do DOCX',1)
    doc.add_paragraph('Itens rastreáveis; aderência estática não equivale a aprovação funcional. Natureza distingue obrigações,recomendações e sínteses repetidas.')
    for section in dict.fromkeys(x['secao'] for x in req):
        doc.add_heading(section,2)
        for x in (r for r in req if r['secao']==section):
            doc.add_heading(x['id']+' | '+x['status']+' | '+x['natureza'],3)
            doc.add_paragraph(x['requisito']);doc.add_paragraph(x['observacao'])
            doc.add_paragraph('Evidência: '+', '.join(x['evidencia_codigo'])+' | Prioridade: '+x['prioridade'])
    doc.add_heading('Apêndice B — Matriz integral das referências visuais',1)
    for x in visual:
        doc.add_heading(x['id']+' | '+x['elemento']+' | '+x['status'],2)
        for k in ['referencia','expectativa','observacao','evidencia']:doc.add_paragraph(k.capitalize()+': '+x[k])
    doc.save(OUT/'RELATORIO_INTERFACE.docx')
    # HTML contains all evidence links and a searchable complete matrix/gallery.
    esc=lambda s:html.escape(str(s),quote=True)
    def html_table(headers,rows,cls=''):
        return '<div class="table-wrap"><table class="'+cls+'"><thead><tr>'+''.join('<th>'+esc(h)+'</th>' for h in headers)+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+esc(c)+'</td>' for c in row)+'</tr>' for row in rows)+'</tbody></table></div>'
    body=[]
    for typ,value in blocks:
        if typ=='h':body.append('<h2>'+esc(value)+'</h2>')
        elif typ=='p':body.append('<p>'+esc(value)+'</p>')
        elif typ=='table':body.append(html_table(*value))
        else:
            file,caption=value;body.append('<figure><a href="capturas/'+esc(file)+'"><img loading="lazy" src="capturas/'+esc(file)+'" alt="'+esc(caption)+'"></a><figcaption>'+esc(caption)+'</figcaption></figure>')
    body.append('<h2 id="matriz">Apêndice A — Matriz integral do DOCX</h2><p>Filtro local por qualquer texto ou estado. Os estados são de inspeção documental/estática; consulte achados e testes para validação de execução.</p><label>Pesquisar requisito <input id="search" type="search" placeholder="Ex.: P0479, lousa, nao_aderente"></label><p id="count"></p>')
    body.append(html_table(['ID / seção / natureza','Requisito','Situação / evidência / observação'],[[x['id']+' | '+x['secao']+' | '+x['natureza'],x['requisito'],x['status']+' | '+', '.join(x['evidencia_codigo'])+' | '+x['observacao']] for x in req],'matrix'))
    body.append('<h2>Apêndice B — Comparação visual integral</h2>'+html_table(['ID / referência','Elemento / esperado','Avaliação / evidência'],[[x['id']+' | '+x['referencia'],x['elemento']+' | '+x['expectativa'],x['status']+' | '+x['observacao']+' | '+x['evidencia']] for x in visual]))
    body.append('<h2>Apêndice C — Controles sem associação de nome identificada</h2>'+html_table(['ID','Linha HTML','Tipo','Placeholder / title'],[[x['id'],x['line'],x['type'],str(x['placeholder'])+' / '+str(x['title'])] for x in stat['controls_without_associated_label']]))
    body.append('<h2 id="galeria">Apêndice D — Galeria de todas as capturas</h2><p>Clique no nome para abrir a imagem na resolução original. Todas são dados fictícios de auditoria.</p><div class="gallery">')
    for path in shots:body.append('<details><summary>'+esc(path.name)+'</summary><a href="capturas/'+esc(path.name)+'"><img loading="lazy" src="capturas/'+esc(path.name)+'" alt="'+esc(path.stem)+'"></a></details>')
    body.append('</div><h2>Apêndice E — Referências originais renderizadas</h2>')
    for path in sorted((OUT/'referencias').glob('*.png')):body.append('<figure><a href="referencias/'+esc(path.name)+'"><img loading="lazy" src="referencias/'+esc(path.name)+'" alt="'+esc(path.stem)+'"></a><figcaption>'+esc(path.name)+'</figcaption></figure>')
    links=['resumo_auditoria.json','testes_fluxos.json','testes_estados.json','testes_integridade_visual.json','metricas_navegador.json','inspecao_estatica.json','manifesto_fontes.json','inventario_referencias.json','cobertura_paragrafos_docx.json','requisitos_extraidos.txt','analise_docx.md','analise_referencias.md','analise_acessibilidade.md']
    body.append('<h2>Apêndice F — Dados e rastreabilidade</h2><ul>'+''.join('<li><a href="dados/'+esc(x)+'">'+esc(x)+'</a></li>' for x in links)+'</ul>')
    css='''body{margin:0;background:#f5f7fa;color:#182331;font:16px/1.55 system-ui,sans-serif}main{max-width:1180px;margin:auto;padding:28px}h2{margin-top:2.3em;color:#123f65;border-bottom:1px solid #c6d3df;padding-bottom:8px}p{overflow-wrap:anywhere}nav{position:sticky;top:0;z-index:2;background:#123f65;color:white;padding:12px 24px;display:flex;gap:20px;flex-wrap:wrap}nav a{color:white}table{border-collapse:collapse;width:100%;font-size:14px;background:white}th,td{border:1px solid #d2dbe4;padding:9px;text-align:left;vertical-align:top;min-width:110px;overflow-wrap:anywhere}th{background:#e5eef5}.table-wrap{overflow:auto;margin:20px 0}figure{background:white;padding:14px;margin:24px 0;border:1px solid #d2dbe4}figure img{max-width:100%;max-height:720px;display:block;margin:auto}figcaption{margin-top:12px;color:#41566a}input{display:block;width:min(95%,650px);padding:12px;font:inherit;margin:12px 0}.gallery{display:grid;grid-template-columns:repeat(auto-fit,minmax(290px,1fr));gap:12px}details{padding:12px;background:white;border:1px solid #d2dbe4;overflow-wrap:anywhere}details img{max-width:100%}a{color:#075e9d}.matrix td:first-child{width:18%}.matrix td:nth-child(2){width:35%}@media print{nav,input,.gallery{display:none}main{max-width:none;padding:0}figure img{max-height:650px}h2{break-after:avoid}tr{break-inside:avoid}body{font-size:10pt;background:white}.table-wrap{overflow:visible}}'''
    js='''const s=document.getElementById('search'),rows=[...document.querySelectorAll('.matrix tbody tr')],count=document.getElementById('count');function filter(){let q=s.value.toLocaleLowerCase('pt-BR');let n=0;for(const r of rows){let show=r.textContent.toLocaleLowerCase('pt-BR').includes(q);r.hidden=!show;if(show)n++}count.textContent=n+' de '+rows.length+' itens exibidos'}s.addEventListener('input',filter);filter();'''
    html_doc='<!doctype html><html lang="pt-BR"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Auditoria completa da interface — 25/10/2026</title><style>'+css+'</style><nav><strong>Auditoria da interface</strong><a href="#matriz">Matriz DOCX</a><a href="#galeria">90 capturas</a><a href="RELATORIO_INTERFACE.docx">Word</a><a href="MATRIZ_REQUISITOS.csv">CSV requisitos</a><a href="ACHADOS_PRIORIZADOS.csv">Achados</a></nav><main>'+''.join(body)+'</main><script>'+js+'</script></html>'
    (OUT/'RELATORIO_INTERFACE.html').write_text(html_doc,encoding='utf-8')
    (OUT/'LEIA-ME.txt').write_text('Abra RELATORIO_INTERFACE.html para consultar o relatório completo,as matrizes e a galeria.\nAlternativa:RELATORIO_INTERFACE.docx.\nAs capturas dependem da pasta capturas e os anexos da pasta dados;preserve toda a estrutura ao copiar.\nNão foram alterados os fontes nem a base do usuário.\nScripts auditam perfis temporários;gerar_relatorio.py consolida dados existentes.\n',encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
